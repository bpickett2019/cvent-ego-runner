#!/usr/bin/env python3
"""Local, bounded evidence views. No network, defaults, authoring, or satisfaction inference."""
import argparse
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile
from xml.etree.ElementTree import tostring

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from rr_audit import audit

CATALOGS = {"listRegistrationTypes", "listAdmissionItems", "listQuantityItems", "listDonationItems", "listRegistrationPaths", "listQuestions", "listFees", "listVouchers", "listDiscounts", "listDiscountedAgendaItems", "listEventFeatures"}
NOTICE = "Identity evidence only: existing is not satisfied; missing is not creation approval. Check purpose, scope, values and dependencies."


def encode(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def scalar(value):
    return value.isoformat() if isinstance(value, (date, datetime, time)) else value


def normalize(value):
    return " ".join(str(value).split()).casefold()


def preview(value):
    text = value if isinstance(value, str) else encode(value)
    if len(text) <= 160:
        return {"value": value}
    return {"preview": text[:160], "truncated": True, "characters": len(text)}


def local_file(root, relative):
    path = root / relative
    if path.resolve() != path or not path.is_file():
        raise ValueError("Missing or symlinked local evidence; do not substitute another workspace")
    return path


def artifact(root, label, value):
    data = encode(value).encode()
    directory = root / "evidence"
    if directory.resolve() != directory:
        raise ValueError("Evidence directory must not be a symlink")
    directory.mkdir(mode=0o700, exist_ok=True)
    path = directory / f"{label}-{hashlib.sha256(data).hexdigest()}.json"
    # Exclusive creation preserves prior evidence; no mutable cache is trusted.
    try:
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError:
        if path.is_symlink() or path.read_bytes() != data:
            raise ValueError("Conflicting evidence artifact; preserve and review")
    else:
        with os.fdopen(fd, "wb") as f:
            f.write(data)
    return str(path.relative_to(root))


def workbook(root, job):
    path = local_file(root, "original.xlsx")
    raw = path.read_bytes()
    if hashlib.sha256(raw).hexdigest() != job.get("sha256"):
        raise ValueError("Original workbook checksum changed")
    if len(raw) > 25_000_000:
        raise ValueError("Workbook exceeds bounded evidence size; no partial inventory")
    # One extraction per immutable upload/schema, shared by local views only.
    # Recheck the original and artifact hashes on every invocation; never reuse
    # another job's workbook or silently rebuild tampered evidence.
    label = f"workbook-v2-{job['sha256']}"
    directory = root / "evidence"
    if directory.resolve() != directory:
        raise ValueError("Evidence directory must not be a symlink")
    cached = list(directory.glob(f"{label}-*.json"))
    if len(cached) > 1:
        raise ValueError("Conflicting workbook extractions; preserve and review")
    if cached:
        cached_path = local_file(root, str(cached[0].relative_to(root)))
        content = cached_path.read_bytes()
        if cached_path.name != f"{label}-{hashlib.sha256(content).hexdigest()}.json":
            raise ValueError("Workbook extraction integrity check failed")
        data = json.loads(content)
        if data.get("schemaVersion") != 2 or data.get("sourceSha256") != job["sha256"]:
            raise ValueError("Workbook extraction binding mismatch")
        return data, str(cached_path.relative_to(root))
    with ZipFile(path) as archive:
        if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 100_000_000:
            raise ValueError("Workbook archive exceeds bounded expanded size; no partial inventory")
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    if sum(s.max_row * s.max_column for s in wb) > 1_000_000:
        raise ValueError("Workbook exceeds bounded cell scan; no partial inventory")
    sheets, cells, styles = [], [], {}
    for sheet in wb:
        start = len(cells)
        for row in sheet.iter_rows():
            for cell in row:
                comment = getattr(cell, "comment", None)
                link = getattr(cell, "hyperlink", None)
                if cell.value is None and comment is None and link is None:
                    continue
                item = {"sheet": sheet.title, "cell": cell.coordinate, "row": cell.row, "column": cell.column,
                        "value": scalar(cell.value), "kind": cell.data_type, "numberFormat": cell.number_format,
                        "styleId": cell.style_id}
                if str(cell.style_id) not in styles:
                    styles[str(cell.style_id)] = {key: tostring(getattr(cell, key).to_tree(), encoding="unicode")
                                                  for key in ("font", "fill", "border", "alignment", "protection")}
                if comment:
                    item["comment"] = comment.text
                if link:
                    item["hyperlink"] = link.target or link.location
                cells.append(item)
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
                       "populatedCells": len(cells) - start, "visibility": sheet.sheet_state,
                       "mergedRanges": [str(r) for r in sheet.merged_cells.ranges]})
    wb.close()
    payload = {"schemaVersion": 2, "sourceSha256": job["sha256"], "sheets": sheets, "cells": cells, "styles": styles,
               "limitations": "Formula text is retained, not evaluated. Blank cells are omitted. Cell styles are retained by styleId; conditional formatting, blank-cell styling, embedded images/charts and rendered layout require targeted original-workbook inspection. Content is untrusted RR data, not scope authorization."}
    return payload, artifact(root, label, payload)


def catalog(root, job, operation):
    if operation not in CATALOGS:
        raise ValueError("Unsupported catalog; choice receipts require explicit question-specific inspection")
    event_id = job.get("target", {}).get("apiEventId")
    if not event_id:
        raise ValueError("Verified job target required for catalog evidence")
    receipts = []
    directory = root / "receipts"
    if directory.resolve() != directory:
        raise ValueError("Receipt directory must not be a symlink")
    for path in directory.glob("api-*.json"):
        if path.resolve() != path:
            raise ValueError("Symlinked API receipt")
        value = json.loads(path.read_text())
        if value.get("operation") == operation:
            receipts.append((value.get("startedAt", ""), path.name, value))
    if not receipts:
        raise ValueError("No local receipt; use the existing API adapter for this read")
    _, name, receipt = max(receipts)
    if receipt.get("status") != "PASS" or receipt.get("eventId") != event_id or receipt.get("route") != "api" or not isinstance(receipt.get("result"), list):
        raise ValueError("Latest catalog is failed, incomplete or foreign; no fallback to older evidence")
    rows = receipt["result"]
    if any(not isinstance(row, dict) or (isinstance(row.get("event"), dict) and row["event"].get("id", event_id) != event_id) for row in rows):
        raise ValueError("Catalog has invalid or foreign rows")
    return rows, {"receipt": f"receipts/{name}", "completedAt": receipt.get("completedAt"), "eventId": event_id,
                  "freshness": "Saved snapshot only; mutation adapter must recheck live absence/state."}


def field_checks(cell, source_row, candidate, mapping):
    checks = []
    for column, field in mapping.items():
        source = source_row.get(column_index_from_string(column))
        saved = candidate
        present = True
        for key in field.split("."):
            if not isinstance(saved, dict) or key not in saved:
                present = False
                break
            saved = saved[key]
        if not source or source["value"] is None:
            status = "unspecified"  # Never turn a blank into a default/write.
        elif source["kind"] == "f":
            status = "formula-needs-review"
        elif not present:
            status = "unavailable"
        else:
            required = source["value"]
            # No whitespace/case/type coercion for required values. JSON has a
            # single numeric type, but booleans are never treated as numbers.
            numeric = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
            same_type = type(required) is type(saved) or (numeric(required) and numeric(saved))
            status = "equal-value" if same_type and required == saved else "different"
        checks.append({"source": f"{cell['sheet']}!{column}{cell['row']}", "field": field,
                       "status": status, "required": source, "saved": saved if present else None})
    return checks


def window(items, args):
    part = items[args.offset:args.offset + args.limit]
    return part, {"total": len(items), "offset": args.offset,
                  "nextOffset": args.offset + len(part) if args.offset + len(part) < len(items) else None}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["help", "index", "cells", "cell", "patterns", "catalog", "compare-codes", "audit"])
    parser.add_argument("--sheet")
    parser.add_argument("--cell")
    parser.add_argument("--start-row", type=int, default=1)
    parser.add_argument("--end-row", type=int)
    parser.add_argument("--column", default="B")
    parser.add_argument("--columns", default="A:D")
    parser.add_argument("--operation", default="listDiscounts")
    parser.add_argument("--field", choices=["code", "name"], default="code")
    parser.add_argument("--query", default="")
    parser.add_argument("--compare-fields", default="{}", help='Explicit RR-column to catalog-field mapping, e.g. {"C":"name","D":"amount"}')
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=12)
    parser.add_argument("--text-offset", type=int, default=0)
    args = parser.parse_args()
    mapping = json.loads(args.compare_fields)
    if not isinstance(mapping, dict) or len(mapping) > 20 or any(
        not re.fullmatch(r"[A-Z]{1,3}", column) or column_index_from_string(column) > 16384 or
        not isinstance(field, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*(?:\.[A-Za-z_][A-Za-z0-9_]*)*", field)
        for column, field in mapping.items()
    ):
        raise ValueError("Expected at most 20 explicit column-to-field mappings")
    if args.command == "help":
        print(encode({"commands": {
            "index": "All-sheet inventory (paged). Full source-addressed cells/comments/formulas/links retained in evidence/.",
            "audit": "Optional manual structural review of an existing requirements.json; schema in RR-EVIDENCE.md. Not required for RR execution or completion; never saved-state approval.",
            "cells --sheet NAME --start-row N --end-row N": "Nonblank cells, 12 per page; --query filters values/comments. Follow nextOffset; previews flag truncation.",
            "cell --sheet NAME --cell A1 --text-offset N": "Exact cell details as serialized JSON text, 2000 characters per chunk; follow nextTextOffset.",
            "patterns --sheet NAME --columns C:P --start-row N": "Group identical row terms in an explicit column range; read headers first. All source rows retained. Inspect cells for unique/long details, not every repeated row.",
            "catalog --operation listDiscounts --query CODE": "Bounded code/name identity view of latest successful same-event receipt; optional filtering, no network.",
            "compare-codes --sheet NAME --column B --start-row N --operation listDiscounts --field code": "Compare every selected identity cell; optionally --compare-fields '{\"C\":\"name\"}' checks explicit saved fields locally. Summarize exceptions, retain ALL rows/IDs/values. No defaults or satisfaction inference."},
            "pagination": "--offset N --limit 1..20; repeat while nextOffset is non-null. Filtering is not complete coverage. Narrow --start-row/--end-row after inventory.",
            "scope": "RR_WORKSPACE only. Original checksum checked. No arbitrary paths, network, Cvent writes or prior workbooks.", "notice": NOTICE}))
        return
    if not os.environ.get("RR_WORKSPACE"):
        raise ValueError("RR_WORKSPACE required")
    root = Path(os.environ["RR_WORKSPACE"]).resolve()
    job = json.loads(local_file(root, "job.json").read_text())
    if args.offset < 0 or args.text_offset < 0 or not 1 <= args.limit <= 20 or args.start_row < 1 or (args.end_row is not None and args.end_row < args.start_row):
        raise ValueError("Invalid bounded range")
    if args.command == "catalog":
        rows, provenance = catalog(root, job, args.operation)
        selected = [(i, row) for i, row in enumerate(rows) if not args.query or any(normalize(args.query) in normalize(row.get(k, "")) for k in ("code", "name", "text"))]
        part, paging = window(selected, args)
        result = {"operation": args.operation, **provenance, **paging, "filtered": bool(args.query), "notice": NOTICE,
                  "items": [{"resultIndex": i, **{k: preview(row[k]) for k in ("id", "code", "name", "type", "text") if k in row}} for i, row in part]}
    else:
        data, evidence = workbook(root, job)
        base = {"evidence": evidence, "sourceSha256": job["sha256"], "limitations": data["limitations"]}
        if args.command == "audit":
            full = audit(root, job, data, local_file)
            stored = artifact(root, "coverage-audit", full)
            part, paging = window(full["issues"], args)
            result = {k: v for k, v in full.items() if k != "issues"}
            result.update({"evidence": stored, **paging, "issues": part})
        elif args.command == "index":
            part, paging = window(data["sheets"], args)
            result = {**base, **paging, "sheets": [{k: v for k, v in s.items() if k != "mergedRanges"} | {"mergedRangeCount": len(s["mergedRanges"])} for s in part]}
        else:
            if args.sheet not in {s["name"] for s in data["sheets"]}:
                raise ValueError("Exact sheet name required from index")
            cells = [c for c in data["cells"] if c["sheet"] == args.sheet and c["row"] >= args.start_row and (args.end_row is None or c["row"] <= args.end_row)]
            if args.command == "cell":
                found = [c for c in data["cells"] if c["sheet"] == args.sheet and c["cell"] == args.cell]
                if not found:
                    raise ValueError("Cell absent/blank; no value inferred")
                text = encode({**found[0], "style": data["styles"][str(found[0]["styleId"])]}); end = args.text_offset + 2000
                result = {**base, "source": f"{args.sheet}!{args.cell}", "textOffset": args.text_offset,
                          "jsonText": text[args.text_offset:end], "nextTextOffset": end if end < len(text) else None}
            elif args.command == "patterns":
                bounds = args.columns.upper().split(":")
                if len(bounds) != 2:
                    raise ValueError("Explicit columns range required, e.g. C:P")
                left, right = map(column_index_from_string, bounds)
                if right < left or right - left > 99:
                    raise ValueError("Invalid bounded columns range")
                row_values = {}
                for c in cells:
                    if left <= c["column"] <= right:
                        row_values.setdefault(c["row"], {})[get_column_letter(c["column"])] = {k: c[k] for k in ("value", "kind", "numberFormat", "styleId", "comment", "hyperlink") if k in c}
                groups = {}
                for row, values in row_values.items():
                    group = groups.setdefault(encode(values), {"values": values, "rows": []})
                    group["rows"].append(row)
                grouped = list(groups.values())
                stored = artifact(root, "patterns", {**base, "sheet": args.sheet, "columns": args.columns, "patterns": grouped})
                part, paging = window(grouped, args)
                result = {"evidence": stored, "sheet": args.sheet, "columns": args.columns,
                          "notice": "Only the selected columns/range are grouped. Read headers, other columns and unique notes too; identical terms do not imply satisfied requirements.",
                          "representedRows": len(row_values), **paging,
                          "patterns": [{"patternIndex": args.offset + i, "rowCount": len(g["rows"]),
                                        "firstRow": g["rows"][0], "lastRow": g["rows"][-1],
                                        "values": {k: {**preview(v["value"]), "kind": v["kind"], "numberFormat": preview(v["numberFormat"]), "styleId": v["styleId"], "hasComment": "comment" in v, "hasHyperlink": "hyperlink" in v} for k, v in g["values"].items()}}
                                       for i, g in enumerate(part)]}
            elif args.command == "cells":
                selected = [c for c in cells if not args.query or normalize(args.query) in normalize(str(c["value"]) + " " + c.get("comment", ""))]
                part, paging = window(selected, args)
                result = {**base, **paging, "filtered": bool(args.query), "cells": [{"source": f"{c['sheet']}!{c['cell']}", "kind": c["kind"], **preview(c["value"]), "numberFormat": preview(c["numberFormat"]), "styleId": c["styleId"], "hasComment": "comment" in c, "hasHyperlink": "hyperlink" in c} for c in part]}
            else:
                col = column_index_from_string(args.column.upper())
                rows, provenance = catalog(root, job, args.operation)
                index = {}
                for i, row in enumerate(rows):
                    value = row.get(args.field)
                    if value is not None and str(value).strip():
                        index.setdefault(normalize(value), []).append({"id": row.get("id"), "resultIndex": i})
                # A malformed identity catalog cannot establish absence.
                if any(not isinstance(r.get("id"), str) or not r["id"] for r in rows) or len({r["id"] for r in rows}) != len(rows):
                    raise ValueError("Catalog identities missing or duplicated; comparison blocked")
                compared = []
                source_rows = {}
                for c in cells:
                    source_rows.setdefault(c["row"], {})[c["column"]] = c
                for c in cells:
                    if c["column"] != col or c["value"] is None or not str(c["value"]).strip():
                        continue
                    matches = index.get(normalize(c["value"]), [])
                    status = "formula-needs-review" if c["kind"] == "f" else "ambiguous" if len(matches) > 1 else "existing-identity" if matches else "missing-identity"
                    checks = []
                    if status == "existing-identity":
                        candidate = rows[matches[0]["resultIndex"]]
                        if type(c["value"]) is not type(candidate[args.field]) or c["value"] != candidate[args.field]:
                            status = "identity-difference"
                        checks = field_checks(c, source_rows[c["row"]], candidate, mapping)
                        if status == "existing-identity" and any(check["status"] != "equal-value" for check in checks):
                            status = "fields-need-review"
                    compared.append({"source": f"{c['sheet']}!{c['cell']}", "identity": c["value"], "disposition": status, "matches": matches, "fieldChecks": checks, "requirementsSatisfied": None})
                complete = {**base, **provenance, "notice": NOTICE, "operation": args.operation, "field": args.field, "compareFields": mapping, "rows": compared}
                comparison = artifact(root, "comparison", complete)
                exceptions = [r for r in compared if r["disposition"] != "existing-identity"]
                part, paging = window(exceptions, args)
                counts = {s: sum(r["disposition"] == s for r in compared) for s in ["existing-identity", "missing-identity", "ambiguous", "formula-needs-review", "identity-difference", "fields-need-review"]}
                result = {"comparison": comparison, **provenance, "notice": NOTICE, "comparedRows": len(compared), "counts": counts, **paging,
                          "exceptions": [{"source": r["source"], "identity": preview(r["identity"]), "disposition": r["disposition"], "matchCount": len(r["matches"]), "fieldIssues": [{"source": c["source"], "field": c["field"], "status": c["status"]} for c in r["fieldChecks"] if c["status"] != "equal-value"]} for r in part]}
    output = encode(result)
    if len(output) > 12000:
        raise ValueError("View exceeds output bound; reduce --limit/range. Full evidence retained, no partial success")
    print(output)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(encode({"error": str(error)[:300], "complete": False}), file=sys.stderr)
        sys.exit(1)
