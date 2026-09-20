#!/usr/bin/env python3
"""Compile the uploaded RR workbook into a concise, source-addressed execution plan."""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PRESERVED_IDENTITIES = ["M-09", "M-10", "M-11", "AGES", "NAICS36D", "CSUB4", "SUB4", "DONATE"]


def scalar(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return re.sub(r"[ \t]+", " ", value).strip()
    return value


def populated_rows(ws):
    for row_no, row in enumerate(ws.iter_rows(), start=1):
        values = [scalar(cell.value) for cell in row]
        if any(value not in (None, "") for value in values):
            yield row_no, values


def adjacent_requirements(ws, start=1, end=None):
    out = []
    for row_no, values in populated_rows(ws):
        if row_no < start or (end and row_no > end):
            continue
        label = values[0] if values else None
        if not isinstance(label, str) or not label:
            continue
        value = values[1] if len(values) > 1 else None
        notes = values[2] if len(values) > 2 else None
        if value not in (None, "") or notes not in (None, ""):
            out.append({"source": f"{ws.title}!A{row_no}", "label": label, "value": value, "notes": notes})
    return out


def parse_questions(ws):
    questions = []
    rows = list(populated_rows(ws))
    starts = []
    for index, (row_no, values) in enumerate(rows):
        identity = values[1] if len(values) > 1 else None
        text = values[3] if len(values) > 3 else None
        if isinstance(identity, str) and identity.strip() and isinstance(text, str) and text.strip():
            starts.append((index, row_no, values))
    for pos, (index, row_no, values) in enumerate(starts):
        next_row = starts[pos + 1][1] if pos + 1 < len(starts) else ws.max_row + 1
        choices = []
        for choice_row in range(row_no + 1, next_row):
            code = scalar(ws.cell(choice_row, 5).value)
            text = scalar(ws.cell(choice_row, 6).value)
            if code not in (None, "") or text not in (None, ""):
                choices.append({"source": f"{ws.title}!E{choice_row}:F{choice_row}", "code": code, "text": text})
        questions.append({
            "source": f"{ws.title}!B{row_no}", "sourceIdentity": values[1], "page": values[0],
            "scope": values[2], "text": values[3], "appearance": values[6] if len(values) > 6 else None,
            "required": values[7] if len(values) > 7 else None, "registrationTypes": values[8] if len(values) > 8 else None,
            "visibleOnline": values[9] if len(values) > 9 else None, "trigger": values[11] if len(values) > 11 else None,
            "notes": values[13] if len(values) > 13 else None, "choices": choices,
            "cventObjectId": None,
        })
    return questions


def parse(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    wb = load_workbook(path, data_only=False, read_only=True)
    sheets = []
    all_cells: list[tuple[str, str, str]] = []
    for ws in wb.worksheets:
        populated = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    populated += 1
                    all_cells.append((ws.title, cell.coordinate, str(cell.value)))
        sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "populatedCells": populated})

    event_ws = wb["Event Details"]
    event_details = adjacent_requirements(event_ws, 9, 42)

    reg_ws = wb["Registration Types & Pricing"]
    registration = []
    admission_items: dict[str, dict[str, Any]] = {}
    for row_no in range(16, 34):
        row = [scalar(reg_ws.cell(row_no, col).value) for col in range(1, 29)]
        if not any(value not in (None, "") for value in row):
            continue
        item = row[1]
        registration.append({
            "source": f"Registration Types & Pricing!A{row_no}:X{row_no}", "registrationType": row[0],
            "admissionItem": item, "displayName": row[2], "badgeDescription": row[3],
            "method": row[4], "reported": row[5], "approvalNeeded": row[6], "guest": row[9],
            "qualified": row[10], "pricing": {"superSaverMember": row[12], "superSaverNonMember": row[13],
            "earlyBirdMember": row[14], "earlyBirdNonMember": row[15], "advanceMember": row[16],
            "advanceNonMember": row[17], "onsiteMember": row[18], "onsiteNonMember": row[19]},
            "description": row[21], "notes": row[23],
        })
        if item:
            admission_items.setdefault(str(item), {"sourceRows": [], "name": item, "usedBy": []})
            admission_items[str(item)]["sourceRows"].append(row_no)
            admission_items[str(item)]["usedBy"].append(row[0])

    mapping_ws = wb["NEW REG MAPPING"]
    reg_mapping = []
    for row_no in range(2, mapping_ws.max_row + 1):
        values = [scalar(mapping_ws.cell(row_no, col).value) for col in range(1, 7)]
        if any(value not in (None, "") for value in values):
            reg_mapping.append({"source": f"NEW REG MAPPING!A{row_no}:F{row_no}", "oldRegistrationType": values[0], "newRegistrationType": values[1], "newCode": values[2], "oldAdmissionItem": values[4], "newAdmissionItem": values[5]})

    policies = adjacent_requirements(wb["Policies & Rules"])
    badge = next((item for item in policies if item["label"].lower() == "badge reprint fee"), None)
    badge_types = next((item for item in policies if item["label"].lower() == "reprint fee reg types"), None)

    optional = []
    session_ws = wb["Sessions_Add-Ons"]
    for row_no in range(4, session_ws.max_row + 1):
        values = [scalar(session_ws.cell(row_no, col).value) for col in range(1, min(32, session_ws.max_column) + 1)]
        code, title = values[0], values[3]
        if code not in (None, "") and title not in (None, ""):
            optional.append({"source": f"Sessions_Add-Ons!A{row_no}:AF{row_no}", "code": code, "sessionboardSync": values[1], "visible": values[2], "title": title, "description": values[4], "allotment": values[5], "startDate": values[6], "startTime": values[7], "endDate": values[8], "endTime": values[9], "pricing": values[10:19], "registrationTypes": values[18] if len(values) > 18 else None})

    discounts = []
    discount_ws = wb["Discount Code Template"]
    for row_no in range(6, discount_ws.max_row + 1):
        values = [scalar(discount_ws.cell(row_no, col).value) for col in range(1, 16)]
        if values[0] not in (None, "") and values[1] not in (None, ""):
            discounts.append({"source": f"Discount Code Template!A{row_no}:O{row_no}", "name": values[0], "code": values[1], "method": values[2], "amount": values[3], "effectiveFrom": values[4], "effectiveTo": values[5], "capacity": values[6], "stackable": values[7], "audience": values[8], "active": values[10], "note": values[11], "admissionItems": values[12], "sessions": values[14]})

    questions = parse_questions(wb["Show Questions"])
    found = {identity: [] for identity in PRESERVED_IDENTITIES}
    for sheet, coordinate, text in all_cells:
        for identity in PRESERVED_IDENTITIES:
            if text.strip().casefold() == identity.casefold():
                found[identity].append(f"{sheet}!{coordinate}")
    identity_map = [{"rrSourceIdentity": key, "sourceLocations": locations, "semanticQuestion": next((q["text"] for q in questions if str(q["sourceIdentity"]).casefold() == key.casefold()), None), "cventObjectId": None, "status": "FOUND" if locations else "NOT_PRESENT_IN_WORKBOOK"} for key, locations in found.items()]

    links = adjacent_requirements(wb["Helpful & Social Media Links"])
    communications = adjacent_requirements(wb["Communications"])
    site_content = adjacent_requirements(wb["Approval Site Parameters"])
    agenda_count = len(optional) + sum(1 for row_no, values in populated_rows(wb["MPT Sessions"]) if row_no > 3 and values and str(values[0] or "").strip())
    speaker_refs = sum(1 for _row_no, values in populated_rows(wb["Reference IDs"]) if any("speaker" in str(value).lower() for value in values if value))

    return {
        "schemaVersion": 1, "sourceFile": path.name, "sourcePath": str(path.resolve()),
        "sourceSha256": hashlib.sha256(raw).hexdigest(), "parsedAt": datetime.now(timezone.utc).isoformat(),
        "workbook": {"sheetCount": len(wb.sheetnames), "sheets": sheets},
        "event": {"requirements": event_details},
        "registration": {"types": registration, "mapping": reg_mapping},
        "admissionItems": list(admission_items.values()),
        "optionalItemsAndSessions": optional,
        "questions": questions,
        "questionIdentityMap": identity_map,
        "discounts": discounts,
        "agenda": {"identifiedRows": agenda_count, "sourceSheets": ["Sessions_Add-Ons", "MPT Sessions"]},
        "speakers": {"identifiedReferenceRows": speaker_refs, "sourceSheet": "Reference IDs"},
        "siteAndContent": {"links": links, "approvalSite": site_content, "communications": communications, "policies": policies},
        "benchmarkTargets": {"eventDetails": event_details, "registrationTypes": registration, "badgeReprintFee": {"fee": badge, "registrationTypes": badge_types}, "questions": [item for item in questions if item["sourceIdentity"] in PRESERVED_IDENTITIES], "siteDesigner": {"links": links[:20]}},
        "warnings": ["M-09/M-10/M-11 were not found in this workbook and are retained as unresolved source identities."] if not any(found[key] for key in ["M-09", "M-10", "M-11"]) else [],
    }


def main() -> None:
    if len(sys.argv) not in (2, 3):
        raise SystemExit("usage: parse_rr.py INPUT.xlsx [OUTPUT.json]")
    source = Path(sys.argv[1])
    output = Path(sys.argv[2]) if len(sys.argv) == 3 else Path("data/current/rr-plan.json")
    plan = parse(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(plan, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({"output": str(output), "sheets": plan["workbook"]["sheetCount"], "registrationTypes": len(plan["registration"]["types"]), "questions": len(plan["questions"]), "discounts": len(plan["discounts"]), "sessions": len(plan["optionalItemsAndSessions"]), "sha256": plan["sourceSha256"]}))


if __name__ == "__main__":
    main()
