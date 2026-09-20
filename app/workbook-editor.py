"""Local RR preview / text-cell editing. Preserve ZIP entries, originals and formulas.
Input/output is JSON on stdin/stdout. Never executes workbook content or formulas.
"""
import sys, json, zipfile, re
from pathlib import PurePosixPath
from xml.dom import minidom
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, coordinate_to_tuple, range_boundaries

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

def checked_archive(path):
    archive = zipfile.ZipFile(path)
    infos = archive.infolist()
    if len(infos) > 5000 or sum(i.file_size for i in infos) > 100_000_000:
        raise ValueError('Workbook exceeds safe expansion limits')
    names = [i.filename for i in infos]
    if len(names) != len(set(names)) or any(n.startswith('/') or '..' in PurePosixPath(n).parts for n in names):
        raise ValueError('Invalid workbook archive')
    if any('vbaProject' in n or n.startswith('_xmlsignatures/') for n in names):
        raise ValueError('Macro-enabled or signed workbooks cannot be edited here')
    for name in names:
        if name.endswith(('.xml', '.rels')):
            xml = archive.read(name).replace(b'\x00', b'').upper()
            if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                raise ValueError('Workbook XML declarations/entities are not supported')
    return archive

def main(req):
    source = req['source']
    with checked_archive(source) as archive:
        wb = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = [{'name': s.title, 'rows': s.max_row or 1, 'columns': s.max_column or 1} for s in wb.worksheets]
            index = req.get('sheet', 0)
            if not isinstance(index, int) or index < 0 or index >= len(sheets):
                raise ValueError('Invalid sheet')
            sheet = wb.worksheets[index]
            if req['action'] == 'preview':
                offset = req.get('offset', 0)
                if not isinstance(offset, int) or offset < 0 or offset >= max(1, sheet.max_row or 1):
                    raise ValueError('Invalid row offset')
                columns = min(sheet.max_column or 1, 64)
                rows = []
                for cells in sheet.iter_rows(min_row=offset + 1, max_row=min(offset + 80, sheet.max_row or 1), max_col=columns):
                    row = []
                    for c in cells:
                        value = '' if c.value is None else str(c.value)
                        row.append({'value': value[:32000], 'formula': c.data_type == 'f'})
                    rows.append(row)
                return {'sheets': sheets, 'sheet': index, 'offset': offset, 'columns': [get_column_letter(i) for i in range(1, columns + 1)], 'rows': rows}
            if req['action'] != 'edit':
                raise ValueError('Unsupported workbook action')
            edits = req['edits']
            if not isinstance(edits, list) or not 1 <= len(edits) <= 500:
                raise ValueError('Save between 1 and 500 cells at once')
            seen = set()
            for edit in edits:
                ref, value = edit.get('cell'), edit.get('value')
                if not isinstance(ref, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,6}', ref) or ref in seen:
                    raise ValueError('Invalid or duplicate cell')
                row, column = coordinate_to_tuple(ref)
                if row > (sheet.max_row or 1) or column > min(sheet.max_column or 1, 64):
                    raise ValueError('Only existing sheet bounds are editable')
                if not isinstance(value, str) or len(value) > 32000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value):
                    raise ValueError('Invalid cell text')
                if sheet[ref].data_type == 'f':
                    raise ValueError('Formula cells are read-only; edit the source workbook for formula changes')
                seen.add(ref)
            workbook = minidom.parseString(archive.read('xl/workbook.xml'))
            sheet_node = workbook.getElementsByTagNameNS(NS, 'sheet')[index]
            relationship_id = sheet_node.getAttributeNS('http://schemas.openxmlformats.org/officeDocument/2006/relationships', 'id')
            rels = minidom.parseString(archive.read('xl/_rels/workbook.xml.rels'))
            matches = [n for n in rels.getElementsByTagName('Relationship') if n.getAttribute('Id') == relationship_id and n.getAttribute('TargetMode') != 'External']
            if len(matches) != 1:
                raise ValueError('Invalid sheet relationship')
            target = matches[0].getAttribute('Target')
            path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            if '..' in PurePosixPath(path).parts or path not in archive.namelist():
                raise ValueError('Unsafe sheet path')
            document = minidom.parseString(archive.read(path))
            for merged in document.getElementsByTagNameNS(NS, 'mergeCell'):
                left, top, right, bottom = range_boundaries(merged.getAttribute('ref'))
                for edit in edits:
                    row, column = coordinate_to_tuple(edit['cell'])
                    if left <= column <= right and top <= row <= bottom and (row, column) != (top, left):
                        raise ValueError('Edit the top-left cell of a merged range only')
            data = document.getElementsByTagNameNS(NS, 'sheetData')[0]
            prefix = (data.prefix + ':') if data.prefix else ''
            for edit in edits:
                ref, value = edit['cell'], edit['value']
                row_num, col_num = coordinate_to_tuple(ref)
                row = next((n for n in data.childNodes if n.nodeType == n.ELEMENT_NODE and n.getAttribute('r') == str(row_num)), None)
                if row is None:
                    row = document.createElementNS(NS, prefix + 'row'); row.setAttribute('r', str(row_num))
                    following = next((n for n in data.childNodes if n.nodeType == n.ELEMENT_NODE and int(n.getAttribute('r')) > row_num), None)
                    data.insertBefore(row, following)
                cell = next((n for n in row.childNodes if n.nodeType == n.ELEMENT_NODE and n.getAttribute('r') == ref), None)
                if cell is None:
                    cell = document.createElementNS(NS, prefix + 'c'); cell.setAttribute('r', ref)
                    following = next((n for n in row.childNodes if n.nodeType == n.ELEMENT_NODE and coordinate_to_tuple(n.getAttribute('r'))[1] > col_num), None)
                    row.insertBefore(cell, following)
                for child in list(cell.childNodes):
                    if child.nodeType == child.ELEMENT_NODE and child.localName in ('f', 'v', 'is'):
                        cell.removeChild(child)
                # Explicit text, never formulas, links, HTML or executable content.
                cell.setAttribute('t', 'inlineStr')
                inline = document.createElementNS(NS, prefix + 'is')
                text = document.createElementNS(NS, prefix + 't'); text.setAttribute('xml:space', 'preserve')
                text.appendChild(document.createTextNode(value)); inline.appendChild(text); cell.appendChild(inline)
            with zipfile.ZipFile(req['destination'], 'x') as output:
                for info in archive.infolist():
                    output.writestr(info, document.toxml(encoding='utf-8') if info.filename == path else archive.read(info.filename))
            return {'editedCells': len(edits)}
        finally:
            wb.close()

try:
    print(json.dumps(main(json.load(sys.stdin))))
except Exception as error:
    print(json.dumps({'error': str(error)[:300]}))
    sys.exit(1)
